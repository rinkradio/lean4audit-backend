from io import BytesIO
import uuid

from fastapi import HTTPException, status, UploadFile
from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session

from app.models.plant import Plant
from app.models.zone import Zone


TEMPLATE_HEADERS = [
    "zone_name",
    "zone_code",
    "description",
    "zone_leader",
    "status",
]


def build_zone_template() -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Zones"

    sheet.append(TEMPLATE_HEADERS)
    sheet.append([
        "Dispatch",
        "DISP-01",
        "Finished goods dispatch area",
        "Rahul Sharma",
        "Active",
    ])

    for cell in sheet[1]:
        cell.font = cell.font.copy(bold=True)

    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 18
    sheet.column_dimensions["C"].width = 36
    sheet.column_dimensions["D"].width = 28
    sheet.column_dimensions["E"].width = 14

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def import_zones_from_excel(
    db: Session,
    file: UploadFile,
    plant_id: uuid.UUID,
):
    plant = (
        db.query(Plant)
        .filter(
            Plant.id == plant_id,
            Plant.is_active.is_(True),
        )
        .first()
    )

    if plant is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Plant not found or inactive.",
        )

    raw = file.file.read()

    try:
        workbook = load_workbook(
            BytesIO(raw),
            read_only=True,
            data_only=True,
        )
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="The uploaded file is not a valid Excel workbook.",
        ) from exc

    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))

    if not rows:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="The Excel file is empty.",
        )

    headers = [
        str(value).strip().lower() if value is not None else ""
        for value in rows[0]
    ]

    required = {"zone_name", "zone_leader"}
    missing = required - set(headers)

    if missing:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Missing required columns: " + ", ".join(sorted(missing)),
        )

    index = {
        header: position
        for position, header in enumerate(headers)
    }

    created = []
    errors = []
    seen_names = set()
    seen_codes = set()

    for row_number, row in enumerate(rows[1:], start=2):
        if not row or not any(value is not None for value in row):
            continue

        def value(column):
            position = index.get(column)
            if position is None or position >= len(row):
                return ""
            return str(row[position]).strip()

        zone_name = value("zone_name")
        zone_code = value("zone_code") or None
        description = value("description") or None
        zone_leader = value("zone_leader")
        status_value = value("status").lower()

        row_errors = []

        if not zone_name:
            row_errors.append("Zone name is required.")

        if not zone_leader:
            row_errors.append("Zone Leader name is required.")

        normalized_name = zone_name.casefold()

        if zone_name and normalized_name in seen_names:
            row_errors.append(
                f"Zone '{zone_name}' is duplicated in the Excel file."
            )
        elif zone_name:
            duplicate_name = (
                db.query(Zone)
                .filter(
                    Zone.plant_id == plant.id,
                    Zone.name.ilike(zone_name),
                )
                .first()
            )
            if duplicate_name:
                row_errors.append(
                    f"Zone '{zone_name}' already exists in this Plant."
                )

        normalized_code = zone_code.casefold() if zone_code else None

        if normalized_code and normalized_code in seen_codes:
            row_errors.append(
                f"Zone code '{zone_code}' is duplicated in the Excel file."
            )
        elif zone_code:
            duplicate_code = (
                db.query(Zone)
                .filter(
                    Zone.plant_id == plant.id,
                    Zone.code.ilike(zone_code),
                )
                .first()
            )
            if duplicate_code:
                row_errors.append(
                    f"Zone code '{zone_code}' already exists in this Plant."
                )

        if row_errors:
            errors.append({
                "row": row_number,
                "message": " ".join(row_errors),
            })
            continue

        is_active = status_value not in {
            "inactive",
            "disabled",
            "false",
            "0",
        }

        zone = Zone(
            name=zone_name,
            code=zone_code,
            description=description,
            plant_id=plant.id,
            zone_leader=zone_leader,
            is_active=is_active,
        )

        db.add(zone)
        created.append(zone)
        seen_names.add(normalized_name)

        if normalized_code:
            seen_codes.add(normalized_code)

    if errors and not created:
        db.rollback()
        return {
            "created_count": 0,
            "error_count": len(errors),
            "errors": errors,
        }

    try:
        db.commit()
    except Exception as exc:
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Unable to import Zones. Please try again.",
        ) from exc

    return {
        "created_count": len(created),
        "error_count": len(errors),
        "errors": errors,
    }
