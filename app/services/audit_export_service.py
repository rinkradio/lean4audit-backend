import io
import os
import uuid
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Image as PdfImage,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
    PageBreak,
)
from sqlalchemy.orm import Session, joinedload

from app.models.audit import Audit, AuditStatus
from app.models.user import User
from app.models.observation import Observation
from app.models.observation_evidence import ObservationEvidence
from app.services.observation_evidence_service import (
    get_storage_path,
)


# =========================================================
# COLORS
# =========================================================

NAVY = "172033"
BLUE = "2563EB"
LIGHT_BLUE = "EFF6FF"
LIGHT_GRAY = "F8FAFC"
MID_GRAY = "E2E8F0"
DARK_GRAY = "475569"
GREEN = "15803D"
RED = "B91C1C"
AMBER = "B45309"
WHITE = "FFFFFF"


# =========================================================
# AUDIT DATA
# =========================================================

def get_export_data(
    db: Session,
    audit_id: uuid.UUID,
):
    audit = (
        db.query(Audit)
        .options(
            joinedload(Audit.zone),
            joinedload(Audit.auditor),
        )
        .filter(Audit.id == audit_id)
        .first()
    )

    if not audit:
        return None

    observations = (
        db.query(Observation)
        .options(
            joinedload(Observation.category),
            joinedload(Observation.creator),
        )
        .filter(
            Observation.audit_id == audit.id
        )
        .order_by(
            Observation.created_at.asc()
        )
        .all()
    )

    evidence_by_observation = {}

    responsible_person_by_id = {}

    responsible_person_ids = [
        observation.responsible_person_id
        for observation in observations
        if observation.responsible_person_id
    ]

    if responsible_person_ids:
        responsible_people = (
            db.query(User)
            .filter(
                User.id.in_(responsible_person_ids)
            )
            .all()
        )

        for person in responsible_people:
            full_name = getattr(
                person,
                "full_name",
                None,
            )

            if not full_name:
                name_parts = [
                    getattr(person, "first_name", None),
                    getattr(person, "last_name", None),
                ]
                full_name = " ".join(
                    part.strip()
                    for part in name_parts
                    if isinstance(part, str) and part.strip()
                )

            responsible_person_by_id[
                person.id
            ] = (
                full_name
                or getattr(person, "email", None)
                or str(person.id)
            )

    if observations:
        observation_ids = [
            observation.id
            for observation in observations
        ]

        evidence_items = (
            db.query(ObservationEvidence)
            .filter(
                ObservationEvidence.observation_id.in_(
                    observation_ids
                )
            )
            .order_by(
                ObservationEvidence.created_at.asc()
            )
            .all()
        )

        for evidence in evidence_items:
            evidence_by_observation.setdefault(
                evidence.observation_id,
                [],
            ).append(evidence)

    return {
        "audit": audit,
        "observations": observations,
        "evidence_by_observation": evidence_by_observation,
        "responsible_person_by_id": responsible_person_by_id,
    }


# =========================================================
# HELPERS
# =========================================================

def _display_date(value):
    if not value:
        return "—"

    if hasattr(value, "strftime"):
        return value.strftime("%d %B %Y")

    return str(value)


def _display_datetime(value):
    if not value:
        return "—"

    if hasattr(value, "strftime"):
        return value.strftime(
            "%d %b %Y, %I:%M %p"
        )

    return str(value)


def _enum_value(value):
    if value is None:
        return "—"

    return getattr(
        value,
        "value",
        str(value),
    )


def _severity_text(value):
    value = _enum_value(value)

    if value == "HIGH":
        return "High"

    if value == "MEDIUM":
        return "Medium"

    if value == "LOW":
        return "Low"

    return value


def _status_text(value):
    value = _enum_value(value)

    if value == "OPEN":
        return "Open"

    if value == "CLOSED":
        return "Closed"

    if value == "SUBMITTED":
        return "Submitted"

    if value == "IN_PROGRESS":
        return "In Progress"

    if value == "DRAFT":
        return "Draft"

    return value


def _observation_type_text(value):
    value = _enum_value(value).upper()
    return {
        "5S": "5S",
        "GEMBA": "Gemba",
        "SAFETY": "Safety",
    }.get(value, value or "5S")


def _observation_type_fill(value):
    value = _enum_value(value).upper()
    return {
        "5S": "DBEAFE",
        "GEMBA": "DCFCE7",
        "SAFETY": "FEF3C7",
    }.get(value, "E2E8F0")


def _observation_type_font(value):
    value = _enum_value(value).upper()
    return {
        "5S": "1D4ED8",
        "GEMBA": GREEN,
        "SAFETY": AMBER,
    }.get(value, DARK_GRAY)


def _observation_details_text(observation):
    details = getattr(observation, "details", None) or {}
    parts = []
    for key, value in details.items():
        if value is None or str(value).strip() == "":
            continue
        label = str(key).replace("_", " ").title()
        parts.append(f"{label}: {value}")
    return "\n".join(parts) if parts else "—"


def _safe_filename(value):
    value = str(value or "audit-report")

    allowed = (
        "abcdefghijklmnopqrstuvwxyz"
        "ABCDEFGHIJKLMNOPQRSTUVWXYZ"
        "0123456789"
        "-_"
    )

    return "".join(
        char if char in allowed else "-"
        for char in value
    )


def _severity_fill(severity):
    severity = _enum_value(severity)

    if severity == "HIGH":
        return "FEE2E2"

    if severity == "MEDIUM":
        return "FEF3C7"

    if severity == "LOW":
        return "E2E8F0"

    return WHITE


def _severity_font(severity):
    severity = _enum_value(severity)

    if severity == "HIGH":
        return RED

    if severity == "MEDIUM":
        return AMBER

    if severity == "LOW":
        return DARK_GRAY

    return NAVY


def _observation_type_key(observation):
    value = _enum_value(getattr(observation, "observation_type", None)).upper()
    return value if value in {"5S", "GEMBA", "SAFETY"} else "5S"


def _type_observations(observations):
    return {
        "5S": [item for item in observations if _observation_type_key(item) == "5S"],
        "GEMBA": [item for item in observations if _observation_type_key(item) == "GEMBA"],
        "SAFETY": [item for item in observations if _observation_type_key(item) == "SAFETY"],
    }


def _add_excel_audit_type_sheet(
    workbook,
    audit,
    audit_type,
    observations,
    evidence_by_observation,
    responsible_person_by_id,
    title_fill,
    header_fill,
    header_font,
    normal_font,
    border,
    center,
    wrap,
    light_fill,
):
    labels = {"5S": "5S AUDIT REPORT", "GEMBA": "GEMBA AUDIT REPORT", "SAFETY": "SAFETY AUDIT REPORT"}
    sheet_names = {"5S": "5S Audit Report", "GEMBA": "Gemba Audit Report", "SAFETY": "Safety Audit Report"}
    sheet = workbook.create_sheet(sheet_names[audit_type])
    label = labels[audit_type]

    sheet.merge_cells("A1:M1")
    sheet["A1"] = f"LEAN4AUDIT — {label}"
    sheet["A1"].fill = title_fill
    sheet["A1"].font = Font(color=WHITE, bold=True, size=18)
    sheet["A1"].alignment = center
    sheet.row_dimensions[1].height = 32

    plant = getattr(getattr(audit.zone, "plant", None), "name", "—")
    sheet.merge_cells("A2:M2")
    sheet["A2"] = (
        f"Audit: {audit.audit_number}  |  Plant: {plant}  |  "
        f"Zone: {getattr(audit.zone, 'name', '—')}  |  Date: {_display_date(audit.audit_date)}"
    )
    sheet["A2"].font = Font(color=DARK_GRAY, bold=True, size=9)
    sheet["A2"].alignment = center

    counts = {
        "Total": len(observations),
        "High": sum(1 for x in observations if _enum_value(x.severity) == "HIGH"),
        "Medium": sum(1 for x in observations if _enum_value(x.severity) == "MEDIUM"),
        "Low": sum(1 for x in observations if _enum_value(x.severity) == "LOW"),
        "Open": sum(1 for x in observations if _enum_value(x.status) == "OPEN"),
        "Closed": sum(1 for x in observations if _enum_value(x.status) == "CLOSED"),
    }

    sheet.merge_cells("A4:M4")
    sheet["A4"] = f"{label} — SUMMARY"
    sheet["A4"].fill = PatternFill("solid", fgColor=BLUE)
    sheet["A4"].font = Font(color=WHITE, bold=True, size=11)

    metric_labels = ["Total Observations", "High", "Medium", "Low", "Open", "Closed"]
    for col, metric in enumerate(metric_labels, start=1):
        sheet.cell(5, col, metric)
        sheet.cell(5, col).fill = PatternFill("solid", fgColor=LIGHT_GRAY)
        sheet.cell(5, col).font = Font(color=DARK_GRAY, bold=True, size=8)
        sheet.cell(5, col).alignment = center
        sheet.cell(5, col).border = border
        sheet.cell(6, col, counts[metric.split()[0]])
        sheet.cell(6, col).font = Font(color=NAVY, bold=True, size=16)
        sheet.cell(6, col).alignment = center
        sheet.cell(6, col).border = border

    headers = [
        "Sr. No.", "Observation No.", "Audit Type", "Category / Area",
        "Location", "Observation / Finding", "Corrective Action", "Severity",
        "Responsible Person", "Target Date", "Status", "Type-specific Details", "Evidence Count",
    ]
    for col, header in enumerate(headers, start=1):
        cell = sheet.cell(8, col, header)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = center

    for row_index, observation in enumerate(observations, start=9):
        responsible = responsible_person_by_id.get(
            observation.responsible_person_id, str(observation.responsible_person_id)
        ) if observation.responsible_person_id else "—"
        evidence_items = evidence_by_observation.get(observation.id, [])
        category = getattr(observation.category, "name", None) or "—"
        values = [
            row_index - 8,
            observation.observation_number,
            _observation_type_text(observation.observation_type),
            category if audit_type == "5S" else (
                (getattr(observation, "details", None) or {}).get("gemba_area_process", "—")
                if audit_type == "GEMBA" else
                (getattr(observation, "details", None) or {}).get("hazard_type", "—")
            ),
            observation.location or "—",
            observation.description or "—",
            observation.corrective_action or "—",
            _severity_text(observation.severity),
            responsible,
            _display_date(observation.target_date),
            _status_text(observation.status),
            _observation_details_text(observation),
            len(evidence_items),
        ]
        for col, value in enumerate(values, start=1):
            cell = sheet.cell(row_index, col, value)
            cell.border = border
            cell.alignment = wrap
            cell.font = normal_font
        type_cell = sheet.cell(row_index, 3)
        type_cell.fill = PatternFill("solid", fgColor=_observation_type_fill(audit_type))
        type_cell.font = Font(color=_observation_type_font(audit_type), bold=True, size=9)
        type_cell.alignment = center
        sev = sheet.cell(row_index, 8)
        sev.fill = PatternFill("solid", fgColor=_severity_fill(observation.severity))
        sev.font = Font(color=_severity_font(observation.severity), bold=True, size=9)
        sev.alignment = center
        status = sheet.cell(row_index, 11)
        status.alignment = center

    widths = [9, 20, 14, 24, 22, 40, 38, 13, 24, 16, 14, 40, 14]
    for col, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    sheet.freeze_panes = "A9"
    last_row = max(8, len(observations) + 8)
    sheet.auto_filter.ref = f"A8:M{last_row}"
    sheet.print_title_rows = "1:8"
    sheet.print_area = f"A1:M{last_row}"
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins.left = 0.25
    sheet.page_margins.right = 0.25
    sheet.page_margins.top = 0.45
    sheet.page_margins.bottom = 0.45
    sheet.oddFooter.left.text = "Lean4Audit • Confidential Audit Report"
    sheet.oddFooter.right.text = "Page &[Page] of &[Pages]"


# =========================================================
# EXCEL
# =========================================================

def build_excel_report(
    db: Session,
    audit_id: uuid.UUID,
):
    data = get_export_data(
        db,
        audit_id,
    )

    if not data:
        return None, None

    audit = data["audit"]
    observations = data["observations"]
    evidence_by_observation = (
        data["evidence_by_observation"]
    )
    responsible_person_by_id = (
        data.get("responsible_person_by_id", {})
    )

    workbook = Workbook()

    summary = workbook.active
    summary.title = "Audit Summary"

    observations_sheet = workbook.create_sheet(
        "Detailed Observations"
    )

    evidence_sheet = workbook.create_sheet(
        "Evidence Register"
    )

    chart_data = workbook.create_sheet(
        "Chart Data"
    )
    chart_data.sheet_state = "hidden"

    # -----------------------------------------------------
    # COMMON STYLES
    # -----------------------------------------------------

    thin_gray = Side(
        style="thin",
        color=MID_GRAY,
    )

    border = Border(
        left=thin_gray,
        right=thin_gray,
        top=thin_gray,
        bottom=thin_gray,
    )

    title_fill = PatternFill(
        "solid",
        fgColor=NAVY,
    )

    section_fill = PatternFill(
        "solid",
        fgColor=BLUE,
    )

    header_fill = PatternFill(
        "solid",
        fgColor="1F4E78",
    )

    light_fill = PatternFill(
        "solid",
        fgColor=LIGHT_GRAY,
    )

    white_font = Font(
        color=WHITE,
        bold=True,
        size=11,
    )

    section_font = Font(
        color=WHITE,
        bold=True,
        size=11,
    )

    header_font = Font(
        color=WHITE,
        bold=True,
        size=10,
    )

    normal_font = Font(
        color=NAVY,
        size=9,
    )

    wrap = Alignment(
        vertical="center",
        wrap_text=True,
    )

    center = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True,
    )

    # -----------------------------------------------------
    # AUDIT SUMMARY HEADER
    # -----------------------------------------------------

    summary.merge_cells("A1:L1")
    summary["A1"] = "LEAN4AUDIT — AUDIT REPORT"
    summary["A1"].fill = title_fill
    summary["A1"].font = Font(
        color=WHITE,
        bold=True,
        size=20,
    )
    summary["A1"].alignment = Alignment(
        horizontal="center",
        vertical="center",
    )
    summary.row_dimensions[1].height = 34

    summary.merge_cells("A2:L2")
    summary["A2"] = (
        f"Audit No. : {audit.audit_number}   |   "
        f"Status : {_status_text(audit.status).upper()}"
    )
    summary["A2"].font = Font(
        color=NAVY,
        bold=True,
        size=11,
    )
    summary["A2"].alignment = Alignment(
        horizontal="center",
        vertical="center",
    )
    summary.row_dimensions[2].height = 23

    summary.merge_cells("A4:L4")
    summary["A4"] = "AUDIT INFORMATION"
    summary["A4"].fill = section_fill
    summary["A4"].font = section_font
    summary["A4"].alignment = Alignment(
        horizontal="left",
        vertical="center",
    )

    plant = getattr(
        getattr(audit.zone, "plant", None),
        "name",
        "—",
    )

    plant_code = getattr(
        getattr(audit.zone, "plant", None),
        "code",
        "—",
    )

    audit_rows = [
        (
            "Audit ID",
            audit.audit_number,
            "Status",
            _status_text(audit.status),
        ),
        (
            "Plant",
            plant,
            "Plant Code",
            plant_code,
        ),
        (
            "Zone",
            getattr(audit.zone, "name", "—"),
            "Audit Date",
            _display_date(audit.audit_date),
        ),
        (
            "Auditor",
            getattr(audit.auditor, "full_name", "—"),
            "Employee ID",
            getattr(audit.auditor, "employee_id", "—"),
        ),
        (
            "Zone Leader",
            audit.zone_leader or "—",
            "HOD",
            audit.hod or "—",
        ),
        (
            "Created",
            _display_datetime(audit.created_at),
            "Submitted",
            _display_datetime(audit.submitted_at),
        ),
    ]

    row = 5

    for label1, value1, label2, value2 in audit_rows:
        summary.cell(row=row, column=1, value=label1)
        summary.cell(row=row, column=2, value=value1)
        summary.cell(row=row, column=7, value=label2)
        summary.cell(row=row, column=8, value=value2)

        summary.merge_cells(
            start_row=row,
            start_column=2,
            end_row=row,
            end_column=6,
        )

        summary.merge_cells(
            start_row=row,
            start_column=8,
            end_row=row,
            end_column=12,
        )

        for column in range(1, 13):
            cell = summary.cell(row=row, column=column)
            cell.border = border
            cell.alignment = wrap
            cell.font = normal_font

            if column in (1, 7):
                cell.fill = light_fill
                cell.font = Font(
                    color=DARK_GRAY,
                    bold=True,
                    size=9,
                )

        row += 1

    # -----------------------------------------------------
    # CALCULATE SUMMARY DATA
    # -----------------------------------------------------

    total = len(observations)

    high = sum(
        1
        for item in observations
        if _enum_value(item.severity) == "HIGH"
    )

    medium = sum(
        1
        for item in observations
        if _enum_value(item.severity) == "MEDIUM"
    )

    low = sum(
        1
        for item in observations
        if _enum_value(item.severity) == "LOW"
    )

    open_count = sum(
        1
        for item in observations
        if _enum_value(item.status) == "OPEN"
    )

    closed_count = sum(
        1
        for item in observations
        if _enum_value(item.status) == "CLOSED"
    )

    type_counts = {
        "5S": sum(1 for item in observations if _enum_value(getattr(item, "observation_type", "5S")) == "5S"),
        "Gemba": sum(1 for item in observations if _enum_value(getattr(item, "observation_type", "5S")) == "GEMBA"),
        "Safety": sum(1 for item in observations if _enum_value(getattr(item, "observation_type", "5S")) == "SAFETY"),
    }

    # -----------------------------------------------------
    # MANAGEMENT METRIC CARDS
    # -----------------------------------------------------

    metric_row = 12

    summary.merge_cells("A11:L11")
    summary["A11"] = "OBSERVATION SUMMARY"
    summary["A11"].fill = section_fill
    summary["A11"].font = section_font
    summary["A11"].alignment = Alignment(
        horizontal="left",
        vertical="center",
    )

    cards = [
        ("TOTAL OBSERVATIONS", total, "D9EAF7"),
        ("OPEN ACTIONS", open_count, "EAF3FF"),
        ("CLOSED ACTIONS", closed_count, "E8F5E9"),
        ("HIGH SEVERITY", high, "FDECEC"),
    ]

    card_starts = [1, 4, 7, 10]

    for (label, value, fill_color), start_col in zip(
        cards,
        card_starts,
    ):
        end_col = start_col + 2

        summary.merge_cells(
            start_row=metric_row,
            start_column=start_col,
            end_row=metric_row,
            end_column=end_col,
        )

        summary.merge_cells(
            start_row=metric_row + 1,
            start_column=start_col,
            end_row=metric_row + 1,
            end_column=end_col,
        )

        label_cell = summary.cell(
            metric_row,
            start_col,
            label,
        )
        value_cell = summary.cell(
            metric_row + 1,
            start_col,
            value,
        )

        fill = PatternFill(
            "solid",
            fgColor=fill_color,
        )

        for current_row in (metric_row, metric_row + 1):
            for current_col in range(
                start_col,
                end_col + 1,
            ):
                cell = summary.cell(
                    current_row,
                    current_col,
                )
                cell.fill = fill
                cell.border = border
                cell.alignment = center

        label_cell.font = Font(
            color=DARK_GRAY,
            bold=True,
            size=8,
        )

        value_cell.font = Font(
            color=NAVY,
            bold=True,
            size=20,
        )

    summary.row_dimensions[metric_row].height = 20
    summary.row_dimensions[metric_row + 1].height = 32

    # -----------------------------------------------------
    # CHART DATA
    # -----------------------------------------------------

    category_rows = [
        ("5S", type_counts["5S"]),
        ("Gemba", type_counts["Gemba"]),
        ("Safety", type_counts["Safety"]),
    ]

    chart_data["A1"] = "Audit Type"
    chart_data["B1"] = "Observations"

    for index, (category, count) in enumerate(
        category_rows,
        start=2,
    ):
        chart_data.cell(index, 1, category)
        chart_data.cell(index, 2, count)

    type_data = [
        ("5S", type_counts["5S"]),
        ("Gemba", type_counts["Gemba"]),
        ("Safety", type_counts["Safety"]),
    ]

    chart_data["J1"] = "Audit Type"
    chart_data["K1"] = "Count"
    for index, (label, count) in enumerate(type_data, start=2):
        chart_data.cell(index, 10, label)
        chart_data.cell(index, 11, count)

    severity_data = [
        ("High", high),
        ("Medium", medium),
        ("Low", low),
    ]

    chart_data["D1"] = "Severity"
    chart_data["E1"] = "Count"

    for index, (label, count) in enumerate(
        severity_data,
        start=2,
    ):
        chart_data.cell(index, 4, label)
        chart_data.cell(index, 5, count)

    status_data = [
        ("Open", open_count),
        ("Closed", closed_count),
    ]

    chart_data["G1"] = "Status"
    chart_data["H1"] = "Count"

    for index, (label, count) in enumerate(
        status_data,
        start=2,
    ):
        chart_data.cell(index, 7, label)
        chart_data.cell(index, 8, count)

    # -----------------------------------------------------
    # CHARTS
    # -----------------------------------------------------

    from openpyxl.chart import (
        BarChart,
        DoughnutChart,
        Reference,
    )
    from openpyxl.chart.label import DataLabelList

    if category_rows:
        category_chart = BarChart()
        category_chart.type = "col"
        category_chart.style = 10
        category_chart.title = "Observations by Audit Type"
        category_chart.y_axis.title = "Observation Count"
        category_chart.x_axis.title = "Audit Type"
        category_chart.height = 7.2
        category_chart.width = 13.5
        category_chart.legend = None

        category_values = Reference(
            chart_data,
            min_col=2,
            min_row=1,
            max_row=len(category_rows) + 1,
        )
        category_labels = Reference(
            chart_data,
            min_col=1,
            min_row=2,
            max_row=len(category_rows) + 1,
        )

        category_chart.add_data(
            category_values,
            titles_from_data=True,
        )
        category_chart.set_categories(category_labels)
        category_chart.varyColors = True
        category_chart.dataLabels = DataLabelList()
        category_chart.dataLabels.showVal = True
        summary.add_chart(
            category_chart,
            "A15",
        )

    severity_chart = DoughnutChart()
    severity_chart.title = "Observations by Severity"
    severity_chart.holeSize = 55
    severity_chart.height = 7.2
    severity_chart.width = 9.5
    severity_chart.legend.position = "r"

    severity_values = Reference(
        chart_data,
        min_col=5,
        min_row=1,
        max_row=4,
    )
    severity_labels = Reference(
        chart_data,
        min_col=4,
        min_row=2,
        max_row=4,
    )

    severity_chart.add_data(
        severity_values,
        titles_from_data=True,
    )
    severity_chart.set_categories(severity_labels)
    severity_chart.dataLabels = DataLabelList()
    severity_chart.dataLabels.showPercent = True
    severity_chart.dataLabels.showLeaderLines = True
    summary.add_chart(
        severity_chart,
        "G15",
    )

    status_chart = DoughnutChart()
    status_chart.title = "Observations by Status"
    status_chart.holeSize = 55
    status_chart.height = 7.2
    status_chart.width = 9.5
    status_chart.legend.position = "r"

    status_values = Reference(
        chart_data,
        min_col=8,
        min_row=1,
        max_row=3,
    )
    status_labels = Reference(
        chart_data,
        min_col=7,
        min_row=2,
        max_row=3,
    )

    status_chart.add_data(
        status_values,
        titles_from_data=True,
    )
    status_chart.set_categories(status_labels)
    status_chart.dataLabels = DataLabelList()
    status_chart.dataLabels.showPercent = True
    status_chart.dataLabels.showLeaderLines = True
    summary.add_chart(
        status_chart,
        "J15",
    )

    # -----------------------------------------------------
    # AUDIT TYPE BREAKDOWN
    # -----------------------------------------------------

    summary.merge_cells("A31:L31")
    summary["A31"] = "AUDIT TYPE BREAKDOWN"
    summary["A31"].fill = section_fill
    summary["A31"].font = section_font
    summary["A31"].alignment = Alignment(horizontal="left", vertical="center")

    type_headers = ["Audit Type", "Observation Count", "Share"]
    for col, value in enumerate(type_headers, start=1):
        cell = summary.cell(row=32, column=col, value=value)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = center

    for row_index, label in enumerate(["5S", "Gemba", "Safety"], start=33):
        count = type_counts[label]
        share = f"{(count / total * 100):.1f}%" if total else "0.0%"
        for col, value in enumerate([label, count, share], start=1):
            cell = summary.cell(row=row_index, column=col, value=value)
            cell.border = border
            cell.alignment = center if col > 1 else wrap
            cell.font = normal_font
        summary.cell(row=row_index, column=1).fill = PatternFill("solid", fgColor=_observation_type_fill(label))
        summary.cell(row=row_index, column=1).font = Font(color=_observation_type_font(label), bold=True, size=9)

    # -----------------------------------------------------
    # DETAILED OBSERVATIONS SHEET
    # -----------------------------------------------------

    observations_sheet.merge_cells("A1:N1")
    observations_sheet["A1"] = (
        "DETAILED OBSERVATIONS — FINDINGS, ACTIONS & PHOTOGRAPHIC EVIDENCE"
    )
    observations_sheet["A1"].fill = title_fill
    observations_sheet["A1"].font = Font(
        color=WHITE,
        bold=True,
        size=16,
    )
    observations_sheet["A1"].alignment = center
    observations_sheet.row_dimensions[1].height = 30

    observations_sheet.merge_cells("A2:N2")
    observations_sheet["A2"] = (
        f"{plant}  |  Zone: {getattr(audit.zone, 'name', '—')}  |  "
        f"Audit: {audit.audit_number}  |  Date: {_display_date(audit.audit_date)}"
    )
    observations_sheet["A2"].font = Font(
        color=DARK_GRAY,
        bold=True,
        size=9,
    )
    observations_sheet["A2"].alignment = center

    headers = [
        "Sr. No.",
        "Evidence 01",
        "Audit Type",
        "5S Category",
        "Location",
        "Observation / Finding",
        "Corrective Action",
        "Severity",
        "Responsible Person",
        "Target Date",
        "Status",
        "Type-specific Details",
        "Evidence 02",
        "Evidence Count",
    ]

    for column, header in enumerate(
        headers,
        start=1,
    ):
        cell = observations_sheet.cell(
            row=4,
            column=column,
            value=header,
        )
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = center

    def add_embedded_image(
        sheet,
        evidence,
        cell_address,
        max_width=120,
        max_height=72,
    ):
        image_path = get_storage_path(evidence)

        if not image_path:
            return False

        image_path = Path(image_path)

        if not image_path.exists():
            return False

        try:
            image = ExcelImage(str(image_path))

            if not image.width or not image.height:
                return False

            ratio = min(
                max_width / image.width,
                max_height / image.height,
                1,
            )

            image.width = max(1, int(image.width * ratio))
            image.height = max(1, int(image.height * ratio))

            sheet.add_image(
                image,
                cell_address,
            )

            return True
        except Exception:
            return False

    for row_index, observation in enumerate(
        observations,
        start=5,
    ):
        evidence_items = evidence_by_observation.get(
            observation.id,
            [],
        )

        category_name = getattr(
            observation.category,
            "name",
            "—",
        ) or "—"

        responsible_person = "—"
        if observation.responsible_person_id:
            responsible_person = responsible_person_by_id.get(
                observation.responsible_person_id,
                str(observation.responsible_person_id),
            )

        values = [
            observation.observation_number,
            "",
            _observation_type_text(getattr(observation, "observation_type", "5S")),
            category_name,
            observation.location or "—",
            observation.description or "—",
            observation.corrective_action or "—",
            _severity_text(observation.severity),
            responsible_person,
            _display_date(observation.target_date),
            _status_text(observation.status),
            _observation_details_text(observation),
            "",
            len(evidence_items),
        ]

        for column, value in enumerate(
            values,
            start=1,
        ):
            cell = observations_sheet.cell(
                row=row_index,
                column=column,
                value=value,
            )
            cell.border = border
            cell.alignment = wrap
            cell.font = normal_font

        severity_cell = observations_sheet.cell(
            row=row_index,
            column=8,
        )
        severity_cell.fill = PatternFill(
            "solid",
            fgColor=_severity_fill(
                observation.severity
            ),
        )
        severity_cell.font = Font(
            color=_severity_font(
                observation.severity
            ),
            bold=True,
            size=9,
        )
        severity_cell.alignment = center

        status_cell = observations_sheet.cell(
            row=row_index,
            column=11,
        )
        status_value = _enum_value(
            observation.status
        )

        if status_value == "CLOSED":
            status_cell.fill = PatternFill(
                "solid",
                fgColor="DCFCE7",
            )
            status_cell.font = Font(
                color=GREEN,
                bold=True,
                size=9,
            )
        else:
            status_cell.fill = PatternFill(
                "solid",
                fgColor="DBEAFE",
            )
            status_cell.font = Font(
                color="1D4ED8",
                bold=True,
                size=9,
            )

        status_cell.alignment = center

        evidence_1 = evidence_items[0] if evidence_items else None
        evidence_2 = evidence_items[1] if len(evidence_items) > 1 else None

        if evidence_1:
            if not add_embedded_image(
                observations_sheet,
                evidence_1,
                f"B{row_index}",
            ):
                observations_sheet.cell(
                    row=row_index,
                    column=2,
                    value="Image unavailable",
                )

        elif not evidence_items:
            observations_sheet.cell(
                row=row_index,
                column=2,
                value="No image",
            )

        if evidence_2:
            if not add_embedded_image(
                observations_sheet,
                evidence_2,
                f"M{row_index}",
            ):
                observations_sheet.cell(
                    row=row_index,
                    column=13,
                    value="Image unavailable",
                )
        else:
            observations_sheet.cell(
                row=row_index,
                column=13,
                value="—",
            )

        observations_sheet.row_dimensions[
            row_index
        ].height = 78

    observation_widths = [
        20, 22, 14, 18, 22, 38, 38, 13, 24, 16, 14, 34, 22, 14,
    ]

    for index, width in enumerate(
        observation_widths,
        start=1,
    ):
        observations_sheet.column_dimensions[
            get_column_letter(index)
        ].width = width

    observations_sheet.freeze_panes = "A5"

    last_observation_row = max(
        4,
        len(observations) + 4,
    )

    observations_sheet.auto_filter.ref = (
        f"A4:N{last_observation_row}"
    )

    observations_sheet.print_title_rows = "1:4"

    # -----------------------------------------------------
    # SEPARATE AUDIT REPORT SHEETS
    # -----------------------------------------------------

    grouped_observations = _type_observations(observations)
    for audit_type in ("5S", "GEMBA", "SAFETY"):
        _add_excel_audit_type_sheet(
            workbook,
            audit,
            audit_type,
            grouped_observations[audit_type],
            evidence_by_observation,
            responsible_person_by_id,
            title_fill,
            header_fill,
            header_font,
            normal_font,
            border,
            center,
            wrap,
            light_fill,
        )

    # -----------------------------------------------------
    # EVIDENCE REGISTER
    # -----------------------------------------------------

    evidence_sheet.merge_cells("A1:G1")
    evidence_sheet["A1"] = (
        "PHOTOGRAPHIC EVIDENCE REGISTER"
    )
    evidence_sheet["A1"].fill = title_fill
    evidence_sheet["A1"].font = Font(
        color=WHITE,
        bold=True,
        size=16,
    )
    evidence_sheet["A1"].alignment = center
    evidence_sheet.row_dimensions[1].height = 30

    evidence_headers = [
        "Observation ID",
        "Audit Type",
        "5S Category",
        "Severity",
        "Image",
        "Original Filename",
        "Uploaded On",
    ]

    for column, header in enumerate(
        evidence_headers,
        start=1,
    ):
        cell = evidence_sheet.cell(
            row=3,
            column=column,
            value=header,
        )
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = center

    evidence_row = 4

    for observation in observations:
        evidence_items = evidence_by_observation.get(
            observation.id,
            [],
        )

        for evidence in evidence_items:
            values = [
                observation.observation_number,
                _observation_type_text(getattr(observation, "observation_type", "5S")),
                getattr(observation.category, "name", "—") or "—",
                _severity_text(observation.severity),
                "",
                evidence.original_filename or "—",
                _display_datetime(evidence.created_at),
            ]

            for column, value in enumerate(
                values,
                start=1,
            ):
                cell = evidence_sheet.cell(
                    row=evidence_row,
                    column=column,
                    value=value,
                )
                cell.border = border
                cell.alignment = wrap
                cell.font = normal_font

            image_path = get_storage_path(evidence)

            if (
                image_path
                and Path(image_path).exists()
            ):
                try:
                    image = ExcelImage(
                        str(image_path)
                    )

                    if image.width and image.height:
                        ratio = min(
                            150 / image.width,
                            100 / image.height,
                            1,
                        )
                        image.width = max(
                            1,
                            int(image.width * ratio),
                        )
                        image.height = max(
                            1,
                            int(image.height * ratio),
                        )

                        evidence_sheet.add_image(
                            image,
                            f"E{evidence_row}",
                        )

                        evidence_sheet.row_dimensions[
                            evidence_row
                        ].height = 85
                    else:
                        raise ValueError(
                            "Invalid image dimensions"
                        )

                except Exception:
                    evidence_sheet.cell(
                        evidence_row,
                        4,
                        "Image unavailable",
                    )
            else:
                evidence_sheet.cell(
                    evidence_row,
                    4,
                    "Image unavailable",
                )

            evidence_row += 1

    evidence_widths = [20, 16, 20, 14, 25, 34, 24]

    for index, width in enumerate(
        evidence_widths,
        start=1,
    ):
        evidence_sheet.column_dimensions[
            get_column_letter(index)
        ].width = width

    evidence_sheet.freeze_panes = "A4"

    if evidence_row > 4:
        evidence_sheet.auto_filter.ref = (
            f"A3:G{evidence_row - 1}"
        )

    evidence_sheet.print_title_rows = "1:3"

    # -----------------------------------------------------
    # SUMMARY COLUMN WIDTHS / FOOTER
    # -----------------------------------------------------

    summary_widths = [
        17,
        13,
        13,
        13,
        13,
        13,
        17,
        13,
        13,
        13,
        13,
        13,
    ]

    for index, width in enumerate(
        summary_widths,
        start=1,
    ):
        summary.column_dimensions[
            get_column_letter(index)
        ].width = width

    summary.merge_cells("A37:F37")
    summary["A37"] = (
        '"Small improvements. A cleaner tomorrow."'
    )
    summary["A32"].font = Font(
        color=DARK_GRAY,
        italic=True,
        size=9,
    )
    summary["A32"].alignment = Alignment(
        horizontal="left",
        vertical="center",
    )

    summary.merge_cells("G37:L37")
    summary["G37"] = (
        f"Generated on: {_display_datetime(datetime.now())}"
    )
    summary["G32"].font = Font(
        color=DARK_GRAY,
        size=9,
    )
    summary["G32"].alignment = Alignment(
        horizontal="right",
        vertical="center",
    )

    # -----------------------------------------------------
    # WORKBOOK / PRINT SETTINGS
    # -----------------------------------------------------

    for sheet in workbook.worksheets:
        sheet.sheet_view.showGridLines = False

        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True

        sheet.page_margins.left = 0.25
        sheet.page_margins.right = 0.25
        sheet.page_margins.top = 0.45
        sheet.page_margins.bottom = 0.45

        sheet.oddFooter.left.text = (
            "Lean4Audit • Confidential Audit Report"
        )
        sheet.oddFooter.center.text = (
            "Audit Report"
        )
        sheet.oddFooter.right.text = (
            "Page &[Page] of &[Pages]"
        )

    summary.freeze_panes = "A5"
    summary.sheet_properties.pageSetUpPr.fitToPage = True
    summary.print_area = "A1:L37"

    observations_sheet.print_area = (
        f"A1:L{last_observation_row}"
    )

    evidence_sheet.print_area = (
        f"A1:F{max(3, evidence_row - 1)}"
    )

    # Active sheet should be the management dashboard.
    workbook.active = 0

    output = io.BytesIO()

    workbook.save(output)

    output.seek(0)

    filename = (
        f"{_safe_filename(audit.audit_number)}"
        f"_Audit_Report_All_Types.xlsx"
    )

    return output, filename


# =========================================================
# PDF
# =========================================================

def build_pdf_report(
    db: Session,
    audit_id: uuid.UUID,
):
    data = get_export_data(
        db,
        audit_id,
    )

    if not data:
        return None, None

    audit = data["audit"]
    observations = data["observations"]
    evidence_by_observation = (
        data["evidence_by_observation"]
    )

    output = io.BytesIO()

    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=14 * mm,
        leftMargin=14 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title=f"Audit Report - {audit.audit_number}",
        author="Lean4Audit",
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "AuditTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=21,
        leading=25,
        textColor=colors.HexColor(
            f"#{NAVY}"
        ),
        alignment=TA_LEFT,
        spaceAfter=4,
    )

    subtitle_style = ParagraphStyle(
        "AuditSubtitle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=13,
        textColor=colors.HexColor(
            f"#{DARK_GRAY}"
        ),
    )

    section_style = ParagraphStyle(
        "Section",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=11,
        leading=14,
        textColor=colors.white,
        spaceBefore=8,
        spaceAfter=7,
    )

    body_style = ParagraphStyle(
        "Body",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=8.5,
        leading=12,
        textColor=colors.HexColor(
            f"#{NAVY}"
        ),
    )

    small_style = ParagraphStyle(
        "Small",
        parent=body_style,
        fontSize=7.5,
        leading=10,
    )

    story = []

    # -----------------------------------------------------
    # COVER HEADER
    # -----------------------------------------------------

    header = Table(
        [
            [
                Paragraph(
                    "LEAN4AUDIT",
                    ParagraphStyle(
                        "Brand",
                        parent=styles["Normal"],
                        fontName="Helvetica-Bold",
                        fontSize=10,
                        textColor=colors.white,
                    ),
                ),
                Paragraph(
                    "AUDIT REPORT",
                    ParagraphStyle(
                        "ReportLabel",
                        parent=styles["Normal"],
                        fontName="Helvetica-Bold",
                        fontSize=10,
                        textColor=colors.white,
                        alignment=TA_CENTER,
                    ),
                ),
                Paragraph(
                    audit.audit_number,
                    ParagraphStyle(
                        "AuditNumber",
                        parent=styles["Normal"],
                        fontName="Helvetica-Bold",
                        fontSize=10,
                        textColor=colors.white,
                        alignment=TA_CENTER,
                    ),
                ),
            ]
        ],
        colWidths=[
            55 * mm,
            70 * mm,
            55 * mm,
        ],
    )

    header.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, -1),
                    colors.HexColor(
                        f"#{NAVY}"
                    ),
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "MIDDLE",
                ),
                (
                    "LEFTPADDING",
                    (0, 0),
                    (-1, -1),
                    8,
                ),
                (
                    "RIGHTPADDING",
                    (0, 0),
                    (-1, -1),
                    8,
                ),
                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    8,
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    8,
                ),
            ]
        )
    )

    story.append(header)
    story.append(Spacer(1, 12))

    story.append(
        Paragraph(
            "Audit Report",
            title_style,
        )
    )

    story.append(
        Paragraph(
            f"Audit Reference: <b>{audit.audit_number}</b>",
            subtitle_style,
        )
    )

    story.append(
        Spacer(1, 10)
    )

    # -----------------------------------------------------
    # AUDIT INFORMATION
    # -----------------------------------------------------

    story.append(
        Table(
            [
                [
                    Paragraph(
                        "AUDIT INFORMATION",
                        section_style,
                    )
                ]
            ],
            colWidths=[180 * mm],
            style=TableStyle(
                [
                    (
                        "BACKGROUND",
                        (0, 0),
                        (-1, -1),
                        colors.HexColor(
                            f"#{BLUE}"
                        ),
                    ),
                    (
                        "LEFTPADDING",
                        (0, 0),
                        (-1, -1),
                        8,
                    ),
                    (
                        "TOPPADDING",
                        (0, 0),
                        (-1, -1),
                        5,
                    ),
                    (
                        "BOTTOMPADDING",
                        (0, 0),
                        (-1, -1),
                        1,
                    ),
                ]
            ),
        )
    )

    plant = getattr(
        getattr(audit.zone, "plant", None),
        "name",
        "—",
    )

    plant_code = getattr(
        getattr(audit.zone, "plant", None),
        "code",
        "—",
    )

    info_data = [
        [
            Paragraph(
                "<b>Audit ID</b>",
                small_style,
            ),
            Paragraph(
                audit.audit_number,
                small_style,
            ),
            Paragraph(
                "<b>Status</b>",
                small_style,
            ),
            Paragraph(
                _status_text(audit.status),
                small_style,
            ),
        ],
        [
            Paragraph(
                "<b>Plant</b>",
                small_style,
            ),
            Paragraph(
                plant,
                small_style,
            ),
            Paragraph(
                "<b>Plant Code</b>",
                small_style,
            ),
            Paragraph(
                plant_code,
                small_style,
            ),
        ],
        [
            Paragraph(
                "<b>Zone</b>",
                small_style,
            ),
            Paragraph(
                audit.zone.name,
                small_style,
            ),
            Paragraph(
                "<b>Audit Date</b>",
                small_style,
            ),
            Paragraph(
                _display_date(
                    audit.audit_date
                ),
                small_style,
            ),
        ],
        [
            Paragraph(
                "<b>Auditor</b>",
                small_style,
            ),
            Paragraph(
                getattr(
                    audit.auditor,
                    "full_name",
                    "—",
                ),
                small_style,
            ),
            Paragraph(
                "<b>Employee ID</b>",
                small_style,
            ),
            Paragraph(
                getattr(
                    audit.auditor,
                    "employee_id",
                    "—",
                ),
                small_style,
            ),
        ],
        [
            Paragraph(
                "<b>Zone Leader</b>",
                small_style,
            ),
            Paragraph(
                audit.zone_leader,
                small_style,
            ),
            Paragraph(
                "<b>HOD</b>",
                small_style,
            ),
            Paragraph(
                audit.hod,
                small_style,
            ),
        ],
    ]

    info_table = Table(
        info_data,
        colWidths=[
            30 * mm,
            60 * mm,
            30 * mm,
            60 * mm,
        ],
    )

    info_table.setStyle(
        TableStyle(
            [
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.5,
                    colors.HexColor(
                        f"#{MID_GRAY}"
                    ),
                ),
                (
                    "BACKGROUND",
                    (0, 0),
                    (0, -1),
                    colors.HexColor(
                        f"#{LIGHT_GRAY}"
                    ),
                ),
                (
                    "BACKGROUND",
                    (2, 0),
                    (2, -1),
                    colors.HexColor(
                        f"#{LIGHT_GRAY}"
                    ),
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "TOP",
                ),
                (
                    "LEFTPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
                (
                    "RIGHTPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
            ]
        )
    )

    story.append(info_table)
    story.append(Spacer(1, 12))

    # -----------------------------------------------------
    # SUMMARY
    # -----------------------------------------------------

    total = len(observations)

    high = sum(
        1
        for item in observations
        if _enum_value(item.severity)
        == "HIGH"
    )

    medium = sum(
        1
        for item in observations
        if _enum_value(item.severity)
        == "MEDIUM"
    )

    low = sum(
        1
        for item in observations
        if _enum_value(item.severity)
        == "LOW"
    )

    open_count = sum(
        1
        for item in observations
        if _enum_value(item.status)
        == "OPEN"
    )

    closed_count = sum(
        1
        for item in observations
        if _enum_value(item.status)
        == "CLOSED"
    )

    type_counts = {
        "5S": sum(1 for item in observations if _enum_value(getattr(item, "observation_type", "5S")) == "5S"),
        "Gemba": sum(1 for item in observations if _enum_value(getattr(item, "observation_type", "5S")) == "GEMBA"),
        "Safety": sum(1 for item in observations if _enum_value(getattr(item, "observation_type", "5S")) == "SAFETY"),
    }

    story.append(
        Table(
            [
                [
                    Paragraph(
                        "AUDIT SUMMARY",
                        section_style,
                    )
                ]
            ],
            colWidths=[180 * mm],
            style=TableStyle(
                [
                    (
                        "BACKGROUND",
                        (0, 0),
                        (-1, -1),
                        colors.HexColor(
                            f"#{BLUE}"
                        ),
                    ),
                    (
                        "LEFTPADDING",
                        (0, 0),
                        (-1, -1),
                        8,
                    ),
                    (
                        "TOPPADDING",
                        (0, 0),
                        (-1, -1),
                        5,
                    ),
                    (
                        "BOTTOMPADDING",
                        (0, 0),
                        (-1, -1),
                        1,
                    ),
                ]
            ),
        )
    )

    summary_data = [
        [
            Paragraph(
                "<b>Total</b>",
                small_style,
            ),
            Paragraph(
                str(total),
                small_style,
            ),
            Paragraph(
                "<b>High</b>",
                small_style,
            ),
            Paragraph(
                str(high),
                small_style,
            ),
            Paragraph(
                "<b>Medium</b>",
                small_style,
            ),
            Paragraph(
                str(medium),
                small_style,
            ),
        ],
        [
            Paragraph(
                "<b>Low</b>",
                small_style,
            ),
            Paragraph(
                str(low),
                small_style,
            ),
            Paragraph(
                "<b>Open</b>",
                small_style,
            ),
            Paragraph(
                str(open_count),
                small_style,
            ),
            Paragraph(
                "<b>Closed</b>",
                small_style,
            ),
            Paragraph(
                str(closed_count),
                small_style,
            ),
        ],
    ]

    summary_table = Table(
        summary_data,
        colWidths=[
            25 * mm,
            20 * mm,
            25 * mm,
            20 * mm,
            25 * mm,
            20 * mm,
        ],
    )

    summary_table.setStyle(
        TableStyle(
            [
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.5,
                    colors.HexColor(
                        f"#{MID_GRAY}"
                    ),
                ),
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, -1),
                    colors.HexColor(
                        f"#{LIGHT_GRAY}"
                    ),
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "MIDDLE",
                ),
                (
                    "ALIGN",
                    (1, 0),
                    (1, -1),
                    "CENTER",
                ),
                (
                    "ALIGN",
                    (3, 0),
                    (3, -1),
                    "CENTER",
                ),
                (
                    "ALIGN",
                    (5, 0),
                    (5, -1),
                    "CENTER",
                ),
                (
                    "LEFTPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
                (
                    "RIGHTPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    7,
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    7,
                ),
            ]
        )
    )

    story.append(summary_table)

    type_summary = Table(
        [
            [
                Paragraph("<b>5S</b>", small_style),
                Paragraph(str(type_counts["5S"]), small_style),
                Paragraph("<b>Gemba</b>", small_style),
                Paragraph(str(type_counts["Gemba"]), small_style),
                Paragraph("<b>Safety</b>", small_style),
                Paragraph(str(type_counts["Safety"]), small_style),
            ]
        ],
        colWidths=[25*mm, 20*mm, 25*mm, 20*mm, 25*mm, 20*mm],
    )
    type_summary.setStyle(TableStyle([
        ("GRID", (0,0), (-1,-1), 0.5, colors.HexColor(f"#{MID_GRAY}")),
        ("BACKGROUND", (0,0), (-1,-1), colors.HexColor(f"#{LIGHT_GRAY}")),
        ("ALIGN", (1,0), (1,0), "CENTER"),
        ("ALIGN", (3,0), (3,0), "CENTER"),
        ("ALIGN", (5,0), (5,0), "CENTER"),
        ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
        ("LEFTPADDING", (0,0), (-1,-1), 6),
        ("RIGHTPADDING", (0,0), (-1,-1), 6),
        ("TOPPADDING", (0,0), (-1,-1), 7),
        ("BOTTOMPADDING", (0,0), (-1,-1), 7),
    ]))
    story.append(Spacer(1, 6))
    story.append(type_summary)
    story.append(PageBreak())

    # -----------------------------------------------------
    # SEPARATE AUDIT TYPE REPORTS
    # -----------------------------------------------------

    grouped_observations = _type_observations(observations)
    type_titles = {
        "5S": "5S AUDIT REPORT",
        "GEMBA": "GEMBA AUDIT REPORT",
        "SAFETY": "SAFETY AUDIT REPORT",
    }

    for type_index, audit_type in enumerate(("5S", "GEMBA", "SAFETY")):
        type_items = grouped_observations[audit_type]
        if type_index > 0:
            story.append(PageBreak())

        story.append(
            Paragraph(
                type_titles[audit_type],
                ParagraphStyle(
                    f"{audit_type}Heading",
                    parent=styles["Heading1"],
                    fontName="Helvetica-Bold",
                    fontSize=18,
                    textColor=colors.HexColor(f"#{_observation_type_font(audit_type)}"),
                    spaceAfter=5,
                ),
            )
        )
        story.append(
            Paragraph(
                f"Audit: <b>{audit.audit_number}</b> &nbsp; | &nbsp; "
                f"Plant: <b>{plant}</b> &nbsp; | &nbsp; "
                f"Observations in this section: <b>{len(type_items)}</b>",
                subtitle_style,
            )
        )
        story.append(Spacer(1, 8))

        if not type_items:
            empty_table = Table(
                [[Paragraph(f"No {type_titles[audit_type].replace(' AUDIT REPORT', '')} observations were recorded for this audit.", body_style)]],
                colWidths=[180 * mm],
            )
            empty_table.setStyle(TableStyle([
                ("BOX", (0,0), (-1,-1), 0.6, colors.HexColor(f"#{MID_GRAY}")),
                ("BACKGROUND", (0,0), (-1,-1), colors.HexColor(f"#{LIGHT_GRAY}")),
                ("LEFTPADDING", (0,0), (-1,-1), 10),
                ("RIGHTPADDING", (0,0), (-1,-1), 10),
                ("TOPPADDING", (0,0), (-1,-1), 12),
                ("BOTTOMPADDING", (0,0), (-1,-1), 12),
            ]))
            story.append(empty_table)
            continue

        for observation in type_items:
            evidence_items = evidence_by_observation.get(observation.id, [])
            heading = Table(
                [[
                    Paragraph(observation.observation_number, ParagraphStyle(
                        f"{audit_type}ObsNumber", parent=body_style, fontName="Helvetica-Bold",
                        fontSize=9, textColor=colors.white)),
                    Paragraph(
                        f"{_observation_type_text(observation.observation_type)} • "
                        f"{getattr(observation.category, 'name', '—') if getattr(observation, 'category', None) else '—'} • "
                        f"{_severity_text(observation.severity)} • {_status_text(observation.status)}",
                        ParagraphStyle(f"{audit_type}ObsMeta", parent=body_style, fontSize=8, textColor=colors.white),
                    )
                ]],
                colWidths=[42 * mm, 138 * mm],
            )
            heading.setStyle(TableStyle([
                ("BACKGROUND", (0,0), (-1,-1), colors.HexColor(f"#{_observation_type_font(audit_type)}")),
                ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
                ("LEFTPADDING", (0,0), (-1,-1), 7),
                ("RIGHTPADDING", (0,0), (-1,-1), 7),
                ("TOPPADDING", (0,0), (-1,-1), 6),
                ("BOTTOMPADDING", (0,0), (-1,-1), 6),
            ]))
            story.append(heading)

            details = getattr(observation, "details", None) or {}
            primary_detail = "—"
            if audit_type == "GEMBA":
                primary_detail = details.get("gemba_area_process") or details.get("what_was_seen") or "—"
            elif audit_type == "SAFETY":
                primary_detail = details.get("hazard_type") or details.get("unsafe_condition_act") or "—"
            else:
                primary_detail = getattr(observation.category, "name", None) or "—"

            observation_info = Table(
                [
                    [Paragraph("<b>Area / Category</b>", small_style), Paragraph(str(primary_detail), small_style), Paragraph("<b>Location</b>", small_style), Paragraph(observation.location or "—", small_style)],
                    [Paragraph("<b>Observation</b>", small_style), Paragraph(observation.description or "—", small_style), Paragraph("<b>Severity</b>", small_style), Paragraph(_severity_text(observation.severity), small_style)],
                    [Paragraph("<b>Type-specific Details</b>", small_style), Paragraph(_observation_details_text(observation), small_style), Paragraph("<b>Status</b>", small_style), Paragraph(_status_text(observation.status), small_style)],
                    [Paragraph("<b>Corrective Action</b>", small_style), Paragraph(observation.corrective_action or "—", small_style), Paragraph("<b>Target Date</b>", small_style), Paragraph(_display_date(observation.target_date), small_style)],
                    [Paragraph("<b>Responsible</b>", small_style), Paragraph(str(observation.responsible_person_id or "—"), small_style), Paragraph("<b>Logged By</b>", small_style), Paragraph(getattr(observation.creator, "full_name", "—"), small_style)],
                ],
                colWidths=[30 * mm, 65 * mm, 30 * mm, 55 * mm],
            )
            observation_info.setStyle(TableStyle([
                ("GRID", (0,0), (-1,-1), 0.5, colors.HexColor(f"#{MID_GRAY}")),
                ("BACKGROUND", (0,0), (0,-1), colors.HexColor(f"#{LIGHT_GRAY}")),
                ("BACKGROUND", (2,0), (2,-1), colors.HexColor(f"#{LIGHT_GRAY}")),
                ("VALIGN", (0,0), (-1,-1), "TOP"),
                ("LEFTPADDING", (0,0), (-1,-1), 5),
                ("RIGHTPADDING", (0,0), (-1,-1), 5),
                ("TOPPADDING", (0,0), (-1,-1), 6),
                ("BOTTOMPADDING", (0,0), (-1,-1), 6),
            ]))
            story.append(observation_info)

            if evidence_items:
                story.append(Spacer(1, 7))
                story.append(Paragraph("<b>Evidence / Images</b>", small_style))
                image_rows = []
                current_row = []
                for evidence in evidence_items:
                    image_path = get_storage_path(evidence)
                    if not image_path or not Path(image_path).exists():
                        continue
                    try:
                        from io import BytesIO
                        from PIL import Image as PILImage, ImageOps
                        with PILImage.open(image_path) as source_image:
                            source_image = ImageOps.exif_transpose(source_image)
                            source_width, source_height = source_image.size
                            if source_width <= 0 or source_height <= 0:
                                continue
                            max_width = 48 * mm
                            max_height = 38 * mm
                            scale = min(max_width / source_width, max_height / source_height, 1.0)
                            image_width = max(1, source_width * scale)
                            image_height = max(1, source_height * scale)
                            normalized = source_image.convert("RGBA")
                            image_buffer = BytesIO()
                            normalized.save(image_buffer, format="PNG", optimize=True)
                            image_buffer.seek(0)
                        current_row.append([
                            PdfImage(image_buffer, width=image_width, height=image_height),
                            Spacer(1, 2),
                            Paragraph(str(evidence.original_filename or "Evidence"), ParagraphStyle(
                                f"{audit_type}ImageName", parent=small_style, alignment=TA_CENTER, fontSize=6.5))
                        ])
                        if len(current_row) == 3:
                            image_rows.append(current_row)
                            current_row = []
                    except Exception:
                        continue
                if current_row:
                    while len(current_row) < 3:
                        current_row.append("")
                    image_rows.append(current_row)
                if image_rows:
                    image_table = Table(image_rows, colWidths=[60 * mm, 60 * mm, 60 * mm])
                    image_table.setStyle(TableStyle([
                        ("VALIGN", (0,0), (-1,-1), "TOP"),
                        ("ALIGN", (0,0), (-1,-1), "CENTER"),
                        ("LEFTPADDING", (0,0), (-1,-1), 4),
                        ("RIGHTPADDING", (0,0), (-1,-1), 4),
                        ("TOPPADDING", (0,0), (-1,-1), 4),
                        ("BOTTOMPADDING", (0,0), (-1,-1), 7),
                    ]))
                    story.append(image_table)
            story.append(Spacer(1, 12))

    # -----------------------------------------------------
    # FOOTER
    # -----------------------------------------------------

    def add_page_number(canvas, document):
        canvas.saveState()

        canvas.setStrokeColor(
            colors.HexColor(
                f"#{MID_GRAY}"
            )
        )

        canvas.line(
            14 * mm,
            11 * mm,
            196 * mm,
            11 * mm,
        )

        canvas.setFont(
            "Helvetica",
            7,
        )

        canvas.setFillColor(
            colors.HexColor(
                f"#{DARK_GRAY}"
            )
        )

        canvas.drawString(
            14 * mm,
            7 * mm,
            "Lean4Audit • Confidential Audit Report",
        )

        canvas.drawRightString(
            196 * mm,
            7 * mm,
            f"Page {document.page}",
        )

        canvas.restoreState()

    document.build(
        story,
        onFirstPage=add_page_number,
        onLaterPages=add_page_number,
    )

    output.seek(0)

    filename = (
        f"{_safe_filename(audit.audit_number)}"
        f"_Audit_Report_All_Types.pdf"
    )

    return output, filename